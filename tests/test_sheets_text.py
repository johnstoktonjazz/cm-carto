"""Сквозной тест текста листов К1 (examples/taganrog/sheets/out/k1.pptx, собирает scripts/build_k1.sh) и подписей
overlay.json карт: по TEXT.md верстки во всем тексте нет прямых кавычек, кавычек без пары, «ё», «№» без
неразрывного пробела перед номером. Замечание Романа: «Детский сад № 20 "Красная Шапочка"» на листе 1."""

import json
import re
from pathlib import Path

import pytest

from cm_carto.render.labels import NBSP

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "examples" / "taganrog" / "sheets" / "out" / "k1.pptx"
IMAGES = ROOT / "examples" / "taganrog" / "maps" / "images"
PAIRS = {"«": "»", "„": "“"}


def problems(text: str) -> list[str]:
    """Нарушения TEXT.md в строке: прямая кавычка, «ё», «№» без неразрывного пробела, кавычка без пары."""
    found = []
    if '"' in text:
        found.append("прямая кавычка")
    if "ё" in text or "Ё" in text:
        found.append("ё")
    if re.search(rf"№(?!{NBSP}\d)", text):
        found.append("№ без неразрывного пробела перед номером")
    stack = []
    for ch in text:
        if ch in PAIRS:
            stack.append(PAIRS[ch])
        elif ch in PAIRS.values():
            if not stack or stack.pop() != ch:
                found.append(f"{ch} без пары")
                break
    if stack:
        found.append(f"{stack[0]} не закрыта")
    return found


@pytest.mark.parametrize("text, ok", [
    (f"Детский сад №{NBSP}20 «Красная Шапочка»", True),
    ("«Парк „Дубки“»", True),
    ('Детский сад № 20 "Красная Шапочка"', False),
    (f"Детский сад №{NBSP}20 „Красная Шапочка", False),
    ("Парк „Дубки“", True),
    ("Дубки“", False),
    ("«Школа", False),
    ("Четвёртая линия", False),
    ("Школа №9", False),
])
def test_problems(text, ok):
    assert (problems(text) == []) is ok


def shapes(group):
    for shape in group:
        yield shape
        if shape.shape_type == 6:  # группа
            yield from shapes(shape.shapes)


def texts(pptx: Path):
    """(где, текст) по всем фигурам всех листов: текст фигур в группах, ячейки таблиц, имена фигур, заметки."""
    from pptx import Presentation

    for n, slide in enumerate(Presentation(pptx).slides, 1):
        for shape in shapes(slide.shapes):
            yield f"лист {n}, имя фигуры", shape.name
            if shape.has_text_frame:
                yield f"лист {n}, {shape.name}", shape.text_frame.text
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield f"лист {n}, {shape.name}, таблица", cell.text
        if slide.has_notes_slide:
            yield f"лист {n}, заметки", slide.notes_slide.notes_text_frame.text


def test_sheets_text():
    pytest.importorskip("pptx")
    if not PPTX.is_file():
        pytest.skip(f"{PPTX} нет: scripts/build_k1.sh")
    seen = list(texts(PPTX))
    assert any("Красная Шапочка" in t for _, t in seen), "подпись сада № 20 пропала с листа 1 — тест ничего не ловит"
    bad = [(where, text, problems(text)) for where, text in seen if problems(text)]
    assert not bad, "\n".join(f"{w}: {t!r} — {', '.join(p)}" for w, t, p in bad)


def test_overlay_text():
    """Подписи, легенда и подписи источников в overlay.json карт К1 — тот же текст, что ляжет на лист."""
    bad = []
    for path in sorted(IMAGES.glob("k1-*.overlay.json")):
        doc = json.loads(path.read_text(encoding="utf-8"))
        for layer in doc["layers"]:
            items = [layer.get("source_label", "")] + [r.get("text", "") for r in layer.get("legend", [])]
            items += [f.get(k, "") for f in layer.get("features", []) for k in ("text", "name")]
            bad += [(path.name, t, problems(t)) for t in items if problems(t)]
    assert not bad, "\n".join(f"{w}: {t!r} — {', '.join(p)}" for w, t, p in bad)
